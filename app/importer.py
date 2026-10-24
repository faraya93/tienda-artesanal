"""Excel bulk import. Column names are the contract a future tool should emit."""

from __future__ import annotations

import os
import sqlite3
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db import CATEGORIES, create_product, replace_media
from app.formatting import parse_price, parse_stock
from app.uploads import MediaError, classify_upload, remove_file, save_bytes

# Exact header row. Do not translate or reorder.
COLUMNS = ["nombre", "descripcion", "precio", "categoria", "stock", "imagenes", "video"]
MAX_ROWS = 500

_COMMENTS = {
    "nombre": "Nombre de la pieza. Obligatorio.",
    "descripcion": "Dos o tres frases. Obligatoria.",
    "precio": "Entero en pesos, sin puntos. Ejemplo: 12990.",
    "categoria": "Exacto: Carteras, Telares, Pulseras u Otros.",
    "stock": "Entero, 0 o más.",
    "imagenes": "Nombres de archivo separados por coma. El primero es la portada. Máximo 6.",
    "video": "Opcional. Un mp4 o webm. Se muestra al final de la galería.",
}


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "productos"
    sheet.append(COLUMNS)
    sheet.append(
        [
            "Cartera de ejemplo",
            "Cartera tejida a mano con lana y algodón. El asa es de cuero vegetal. Reemplaza esta fila por tus piezas.",
            19990,
            "Carteras",
            3,
            "foto1.jpg, foto2.jpg",
            "",
        ]
    )
    for index, name in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B5541E")
        cell.comment = Comment(_COMMENTS[name], "Taller")
        sheet.column_dimensions[get_column_letter(index)].width = 28 if index != 2 else 72
    sheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def import_workbook(connection, spreadsheet: bytes, media_files: dict[str, tuple[str, bytes]]) -> dict:
    """Create one product per valid row. Invalid rows are reported and not inserted.

    media_files maps a casefolded basename to (original name, bytes).
    """
    report = {"file_error": None, "rows": [], "created": 0, "skipped": 0, "errors": 0}
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(spreadsheet), data_only=False)
    except Exception:
        report["file_error"] = "No pude leer el Excel. Usa la planilla .xlsx de esta página."
        return report

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        report["file_error"] = "La planilla está vacía."
        return report

    header = _header(rows[0])
    if header != COLUMNS:
        report["file_error"] = (
            "La primera fila debe ser exactamente: " + ", ".join(COLUMNS) + "."
        )
        return report

    data_rows = rows[1:]
    while data_rows and _empty(data_rows[-1]):
        data_rows.pop()
    if len(data_rows) > MAX_ROWS:
        report["file_error"] = f"La planilla tiene más de {MAX_ROWS} filas."
        return report

    for offset, row in enumerate(data_rows):
        excel_row = offset + 2
        if _empty(row):
            report["rows"].append(
                {"row": excel_row, "name": "", "result": "skipped", "detail": "Fila vacía."}
            )
            report["skipped"] += 1
            continue
        outcome = _import_row(connection, row, media_files)
        outcome["row"] = excel_row
        report["rows"].append(outcome)
        if outcome["result"] == "created":
            report["created"] += 1
            connection.commit()
        else:
            connection.rollback()
            report["errors"] += 1
    return report


def read_uploaded_media(storages) -> tuple[dict[str, tuple[str, bytes]] | None, str | None]:
    files: dict[str, tuple[str, bytes]] = {}
    for storage in storages:
        original = os.path.basename((storage.filename or "").strip())
        if not original:
            continue
        if original != (storage.filename or "").strip() or ".." in original.split("/"):
            return None, "Un archivo trae una ruta. Sube solo el nombre del archivo."
        key = original.casefold()
        if key in files:
            return None, f"Hay dos archivos con el mismo nombre: {original}."
        files[key] = (original, storage.read())
    return files, None


def _import_row(connection, row, media_files: dict[str, tuple[str, bytes]]) -> dict:
    cells = list(row or [])
    while len(cells) < len(COLUMNS):
        cells.append(None)
    # 0 nombre, 1 descripcion, 2 precio, 3 categoria, 4 stock, 5 imagenes, 6 video
    name = _text(cells[0])
    description = _text(cells[1])
    category = _text(cells[3])
    saved: list[str] = []
    try:
        if not name or len(name) > 80:
            raise ValueError("El nombre es obligatorio y puede tener hasta 80 caracteres.")
        if not description or len(description) > 2000:
            raise ValueError("La descripción es obligatoria y puede tener hasta 2000 caracteres.")
        price = parse_price(cells[2])
        if category not in CATEGORIES:
            raise ValueError(
                "Categoría desconocida. Usa Carteras, Telares, Pulseras u Otros."
            )
        stock = parse_stock(cells[4])
        image_names = _filenames(cells[5])
        video_name = _text(cells[6])
        if not image_names:
            raise ValueError("Falta al menos una imagen en la columna imagenes.")
        if len(image_names) > 6:
            raise ValueError("Puedes indicar hasta 6 imágenes.")
        images = [_take_file(media_files, filename, "image") for filename in image_names]
        videos = []
        if video_name:
            videos.append(_take_file(media_files, video_name, "video"))
        product_id = create_product(
            connection,
            name=name,
            description=description,
            price_clp=price,
            category=category,
            stock=stock,
            active=1,
        )
        items = []
        for media_type, ext, data in images + videos:
            filename = save_bytes(data, ext)
            saved.append(filename)
            items.append({"filename": filename, "media_type": media_type})
        replace_media(connection, product_id, items)
    except (ValueError, MediaError, OSError, sqlite3.IntegrityError) as exc:
        for filename in saved:
            remove_file(filename)
        return {"name": name, "result": "error", "detail": str(exc)}
    return {"name": name, "result": "created", "detail": f"Pieza #{product_id}"}


def _take_file(media_files, filename: str, expected: str):
    key = filename.casefold()
    found = media_files.get(key)
    if found is None:
        raise ValueError(f"No subiste el archivo {filename}.")
    original, data = found
    media_type, ext = classify_upload(original, data)
    if media_type != expected:
        if expected == "image":
            raise ValueError(f"{filename} no es una imagen jpg, png o webp.")
        raise ValueError(f"{filename} no es un video mp4 o webm.")
    return media_type, ext, data


def _filenames(value: object) -> list[str]:
    text = "" if value is None else str(value).strip()
    if not text:
        return []
    names = []
    for part in text.split(","):
        name = os.path.basename(part.strip())
        if not name:
            continue
        if name != part.strip() or ".." in name or "/" in part or "\\" in part:
            raise ValueError(f"Nombre de archivo inválido: {part.strip()}.")
        names.append(name)
    return names


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header(row) -> list[str]:
    cells = list(row or [])
    while cells and (cells[-1] is None or str(cells[-1]).strip() == ""):
        cells.pop()
    return ["" if cell is None else str(cell).strip() for cell in cells]


def _empty(row) -> bool:
    for cell in row or []:
        if cell is not None and str(cell).strip() != "":
            return False
    return True

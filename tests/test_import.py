import io

from openpyxl import Workbook, load_workbook

from app.db import get_db, list_media
from app.importer import COLUMNS
from tests.conftest import jpeg_bytes, png_bytes

WEBM = b"\x1a\x45\xdf\xa3" + b"\x00" * 16


def test_template_has_the_contract_columns(admin):
    response = admin["client"].get("/admin/importar/plantilla")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    header = [cell.value for cell in next(workbook.active.iter_rows(min_row=1, max_row=1))]
    assert header == COLUMNS
    sample = [cell.value for cell in next(workbook.active.iter_rows(min_row=2, max_row=2))]
    assert sample[0] == "Cartera de ejemplo"
    assert sample[5] == "foto1.jpg, foto2.jpg"


def test_import_creates_valid_rows_and_reports_invalid_ones(admin, app):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(COLUMNS)
    sheet.append(
        [
            "Bolso importado",
            "Bolso de lona teñida a mano. El asa es de cuero.",
            24990,
            "Carteras",
            4,
            "foto-a.png, foto-b.jpg",
            "clip.webm",
        ]
    )
    sheet.append([None, None, None, None, None, None, None])
    sheet.append(
        [
            "Pieza inválida",
            "Esta fila no debe crearse.",
            10000,
            "Zapatos",
            1,
            "foto-a.png",
            "",
        ]
    )
    payload = io.BytesIO()
    workbook.save(payload)
    response = admin["client"].post(
        "/admin/importar",
        data={
            "csrf_token": admin["token"],
            "planilla": (io.BytesIO(payload.getvalue()), "productos.xlsx"),
            "archivos": [
                (io.BytesIO(png_bytes()), "foto-a.png"),
                (io.BytesIO(jpeg_bytes()), "foto-b.jpg"),
                (io.BytesIO(WEBM), "clip.webm"),
            ],
        },
    )
    html = response.get_data(as_text=True)
    assert response.status_code == 200
    assert 'data-result="created"' in html
    assert 'data-result="skipped"' in html
    assert 'data-result="error"' in html
    assert "Zapatos" in html or "Categoría desconocida" in html

    with app.app_context():
        names = [row["name"] for row in get_db().execute("SELECT name FROM products")]
        assert names == ["Bolso importado"]
        product_id = get_db().execute("SELECT id FROM products").fetchone()["id"]
        media = list_media(get_db(), product_id)
        assert [row["media_type"] for row in media] == ["image", "image", "video"]
        assert [row["position"] for row in media] == [0, 1, 2]
        assert media[0]["filename"].endswith(".png")
        assert media[1]["filename"].endswith(".jpg")
        assert media[2]["filename"].endswith(".webm")
        assert "Pieza inválida" not in names
